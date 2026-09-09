from docx import Document
from docx.shared import Pt, RGBColor
from io import BytesIO
from fpdf import FPDF


def _wrap_long_words(text: str, max_word_length: int = 40) -> str:
    """
    Inserts a space after every 'max_word_length' characters inside any
    single unbroken word (e.g. a very long URL or run-on term), so the
    PDF renderer can always wrap the line. Without this, a single very
    long word with no spaces can crash PDF generation.
    """
    if not text:
        return text

    words = text.split(" ")
    fixed_words = []

    for word in words:
        if len(word) > max_word_length:
            chunks = [word[i:i + max_word_length] for i in range(0, len(word), max_word_length)]
            fixed_words.append(" ".join(chunks))
        else:
            fixed_words.append(word)

    return " ".join(fixed_words)


def _sanitize_pdf_text(text: str) -> str:
    """
    Ensures text only contains safe, standard ASCII characters that the
    built-in PDF fonts (Helvetica, Times) can always render and measure
    correctly, and that no single unbroken word is too long to wrap.
    """
    if not text:
        return text

    common_replacements = {
        "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-",
        "\u2026": "...", "\u2022": "-",
    }
    for bad_char, replacement in common_replacements.items():
        text = text.replace(bad_char, replacement)

    text = text.encode("ascii", errors="ignore").decode("ascii")
    text = _wrap_long_words(text)

    return text


# ---------------------------------------------------------------------------
# STYLE CONFIGURATION
# ---------------------------------------------------------------------------

TEMPLATE_STYLES = {
    "ats_friendly": {
        "docx_font": "Arial",
        "docx_heading_color": RGBColor(0x00, 0x00, 0x00),
        "pdf_font": "Helvetica",
        "pdf_heading_color": (0, 0, 0),
        "heading_underline": False,
    },
    "modern": {
        "docx_font": "Calibri",
        "docx_heading_color": RGBColor(0x1F, 0x4E, 0x79),
        "pdf_font": "Helvetica",
        "pdf_heading_color": (31, 78, 121),
        "heading_underline": False,
    },
    "professional": {
        "docx_font": "Georgia",
        "docx_heading_color": RGBColor(0x00, 0x00, 0x00),
        "pdf_font": "Times",
        "pdf_heading_color": (0, 0, 0),
        "heading_underline": True,
    },
}


def _set_default_font(document, font_name, size=11):
    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size)


def _add_docx_heading(document, text, style_config):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = style_config["docx_heading_color"]
    run.font.name = style_config["docx_font"]

    if style_config["heading_underline"]:
        run.underline = True

    return paragraph


def generate_docx_resume(tailored: dict, template: str = "ats_friendly") -> BytesIO:
    style_config = TEMPLATE_STYLES.get(template, TEMPLATE_STYLES["ats_friendly"])

    document = Document()
    _set_default_font(document, style_config["docx_font"])

    if tailored.get("summary"):
        _add_docx_heading(document, "Professional Summary", style_config)
        document.add_paragraph(tailored["summary"])

    if tailored.get("skills"):
        _add_docx_heading(document, "Skills", style_config)
        document.add_paragraph(", ".join(tailored["skills"]))

    if tailored.get("experience"):
        _add_docx_heading(document, "Experience", style_config)
        for job in tailored["experience"]:
            heading_text = f"{job.get('title', '')} — {job.get('company', '')}"
            if job.get("duration"):
                heading_text += f" ({job['duration']})"

            p = document.add_paragraph()
            run = p.add_run(heading_text)
            run.bold = True

            for bullet in job.get("bullets", []):
                document.add_paragraph(bullet, style="List Bullet")

    if tailored.get("education"):
        _add_docx_heading(document, "Education", style_config)
        for edu in tailored["education"]:
            heading_text = f"{edu.get('degree', '')} — {edu.get('institution', '')}"
            if edu.get("duration"):
                heading_text += f" ({edu['duration']})"

            p = document.add_paragraph()
            run = p.add_run(heading_text)
            run.bold = True

    if tailored.get("projects"):
        _add_docx_heading(document, "Projects", style_config)
        for project in tailored["projects"]:
            p = document.add_paragraph()
            run = p.add_run(project.get("name", ""))
            run.bold = True

            if project.get("description"):
                document.add_paragraph(project["description"])

            if project.get("technologies"):
                document.add_paragraph("Technologies: " + ", ".join(project["technologies"]))

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _add_pdf_heading(pdf, text, style_config):
    """
    Adds a styled section heading to a PDF, always starting from the
    left margin, using the font and color defined in this template's
    style config.
    """
    text = _sanitize_pdf_text(text)
    pdf.set_x(pdf.l_margin)
    pdf.set_font(style_config["pdf_font"], style="B", size=13)
    pdf.set_text_color(*style_config["pdf_heading_color"])
    try:
        pdf.cell(0, 10, text, new_x="LMARGIN", new_y="NEXT")
    except Exception:
        pdf.set_x(pdf.l_margin)
        pdf.cell(0, 10, "[Heading]", new_x="LMARGIN", new_y="NEXT")

    if style_config["heading_underline"]:
        y = pdf.get_y() - 2
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)

    pdf.set_text_color(0, 0, 0)
    pdf.set_font(style_config["pdf_font"], size=11)


def generate_pdf_resume(tailored: dict, template: str = "ats_friendly") -> BytesIO:
    """
    Builds a PDF resume from the tailored resume data, styled according
    to the chosen template. Returns the file as an in-memory buffer,
    ready for download.
    """
    style_config = TEMPLATE_STYLES.get(template, TEMPLATE_STYLES["ats_friendly"])
    font = style_config["pdf_font"]

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font(font, size=11)

    def safe_multi_cell(text):
        """
        Writes text to the PDF, always starting from the left margin,
        and if rendering still somehow fails, skips that specific piece
        of text instead of crashing the entire PDF generation.
        """
        pdf.set_x(pdf.l_margin)
        try:
            pdf.multi_cell(0, 7, text)
        except Exception:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 7, "[Content could not be displayed]")

    def add_paragraph(text):
        safe_multi_cell(_sanitize_pdf_text(text))
        pdf.ln(2)

    def add_bullet(text):
        safe_multi_cell(f"-  {_sanitize_pdf_text(text)}")

    if tailored.get("summary"):
        _add_pdf_heading(pdf, "Professional Summary", style_config)
        add_paragraph(tailored["summary"])

    if tailored.get("skills"):
        _add_pdf_heading(pdf, "Skills", style_config)
        add_paragraph(", ".join(tailored["skills"]))

    if tailored.get("experience"):
        _add_pdf_heading(pdf, "Experience", style_config)
        for job in tailored["experience"]:
            heading_text = f"{job.get('title', '')} - {job.get('company', '')}"
            if job.get("duration"):
                heading_text += f" ({job['duration']})"
            pdf.set_font(font, style="B", size=11)
            safe_multi_cell(_sanitize_pdf_text(heading_text))
            pdf.set_font(font, size=11)

            for bullet in job.get("bullets", []):
                add_bullet(bullet)
            pdf.ln(2)

    if tailored.get("education"):
        _add_pdf_heading(pdf, "Education", style_config)
        for edu in tailored["education"]:
            heading_text = f"{edu.get('degree', '')} - {edu.get('institution', '')}"
            if edu.get("duration"):
                heading_text += f" ({edu['duration']})"
            safe_multi_cell(_sanitize_pdf_text(heading_text))
        pdf.ln(2)

    if tailored.get("projects"):
        _add_pdf_heading(pdf, "Projects", style_config)
        for project in tailored["projects"]:
            pdf.set_font(font, style="B", size=11)
            safe_multi_cell(_sanitize_pdf_text(project.get("name", "")))
            pdf.set_font(font, size=11)

            if project.get("description"):
                add_paragraph(project["description"])

            if project.get("technologies"):
                add_paragraph("Technologies: " + ", ".join(project["technologies"]))

    buffer = BytesIO(pdf.output())
    buffer.seek(0)
    return buffer